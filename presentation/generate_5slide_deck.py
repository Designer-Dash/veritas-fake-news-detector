import os
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def create_veritas_5slide_presentation(output_path="Veritas_Presentation.pptx"):
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Harmonious Executive Dark Palette
    BG_COLOR = RGBColor(15, 23, 42)          # Deep Navy/Slate #0F172A
    CARD_BG = RGBColor(30, 41, 59)           # Slate Card #1E293B
    CARD_BORDER = RGBColor(51, 65, 85)       # Border #334155
    CYAN = RGBColor(56, 189, 248)            # Accent Cyan #38BDF8
    EMERALD = RGBColor(52, 211, 153)         # WORK DONE / VERIFIED #34D399
    AMBER = RGBColor(251, 191, 36)           # WORK NOT DONE / PENDING #FBBF24
    ROSE = RGBColor(248, 113, 113)           # Danger / Limitations #F87171
    PURPLE = RGBColor(192, 132, 252)         # AI & Data #C084FC
    TEXT_WHITE = RGBColor(248, 250, 252)     # Crisp White #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184)     # Silver #94A3B8
    TEXT_DIM = RGBColor(100, 116, 139)       # Slate Gray #64748B

    TOTAL_SLIDES = 5

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title, category="VERITAS", badge=None, badge_color=CYAN, slide_num=None):
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(10.5), Inches(0.35))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        
        run_cat = p_tag.add_run()
        run_cat.text = category.upper() + ("  •  " if badge else "")
        run_cat.font.size = Pt(10)
        run_cat.font.bold = True
        run_cat.font.color.rgb = TEXT_DIM

        if badge:
            run_badge = p_tag.add_run()
            run_badge.text = f"[{badge.upper()}]"
            run_badge.font.size = Pt(10.5)
            run_badge.font.bold = True
            run_badge.font.color.rgb = badge_color

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(10.5), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.45), Inches(11.733), Inches(0.02)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = CARD_BORDER
        line.line.fill.background()

        if slide_num:
            num_box = slide.shapes.add_textbox(Inches(11.3), Inches(0.4), Inches(1.2), Inches(0.4))
            p_num = num_box.text_frame.paragraphs[0]
            p_num.text = f"{slide_num:02d} / {TOTAL_SLIDES:02d}"
            p_num.alignment = PP_ALIGN.RIGHT
            p_num.font.size = Pt(11)
            p_num.font.bold = True
            p_num.font.color.rgb = TEXT_DIM

    def add_card(slide, left, top, width, height, title="", items=None, accent=CYAN, bg=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = bg
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)

        if accent:
            stripe = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(left + 0.15), Inches(top + 0.02), Inches(width - 0.3), Inches(0.04)
            )
            stripe.fill.solid()
            stripe.fill.fore_color.rgb = accent
            stripe.line.fill.background()

        tb = slide.shapes.add_textbox(
            Inches(left + 0.22), Inches(top + 0.15), Inches(width - 0.44), Inches(height - 0.3)
        )
        tf = tb.text_frame
        tf.word_wrap = True

        if title:
            p_title = tf.paragraphs[0]
            p_title.text = title
            p_title.font.size = Pt(13.5)
            p_title.font.bold = True
            p_title.font.color.rgb = accent or TEXT_WHITE
            p_title.space_after = Pt(6)

        if items:
            for idx, item in enumerate(items):
                p = tf.add_paragraph() if (title or idx > 0) else tf.paragraphs[0]
                if isinstance(item, tuple):
                    bold_prefix, regular_text = item
                    
                    if bold_prefix.startswith("---"):
                        run_sep = p.add_run()
                        run_sep.text = "⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯"
                        run_sep.font.size = Pt(8)
                        run_sep.font.color.rgb = CARD_BORDER
                        p.space_after = Pt(2)
                        continue

                    run1 = p.add_run()
                    run1.text = bold_prefix + " "
                    run1.font.bold = True
                    run1.font.size = Pt(10)
                    
                    if "[DONE]" in bold_prefix:
                        run1.font.color.rgb = EMERALD
                    elif "[NOT DONE]" in bold_prefix or "[PLANNED]" in bold_prefix:
                        run1.font.color.rgb = AMBER
                    elif "[LIMITATION]" in bold_prefix or "[GAP]" in bold_prefix:
                        run1.font.color.rgb = ROSE
                    else:
                        run1.font.color.rgb = TEXT_WHITE

                    run2 = p.add_run()
                    run2.text = regular_text
                    run2.font.bold = False
                    run2.font.size = Pt(9.5)
                    run2.font.color.rgb = TEXT_MUTED
                else:
                    run = p.add_run()
                    run.text = item
                    run.font.size = Pt(10)
                    run.font.color.rgb = TEXT_MUTED
                p.space_after = Pt(3.5)
        return card

    def add_stat_box(slide, left, top, width, height, number, label, sublabel="", accent=CYAN):
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = CARD_BORDER
        box.line.width = Pt(1)

        tb = slide.shapes.add_textbox(
            Inches(left + 0.1), Inches(top + 0.1), Inches(width - 0.2), Inches(height - 0.2)
        )
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = number
        p1.alignment = PP_ALIGN.CENTER
        p1.font.size = Pt(25)
        p1.font.bold = True
        p1.font.color.rgb = accent

        p2 = tf.add_paragraph()
        p2.text = label
        p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE
        p2.space_before = Pt(2)

        if sublabel:
            p3 = tf.add_paragraph()
            p3.text = sublabel
            p3.alignment = PP_ALIGN.CENTER
            p3.font.size = Pt(8.5)
            p3.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 1: TITLE SLIDE & EXECUTIVE SUMMARY
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    glow = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.12), Inches(3.6))
    glow.fill.solid()
    glow.fill.fore_color.rgb = CYAN
    glow.line.fill.background()

    t_box = s1.shapes.add_textbox(Inches(1.2), Inches(1.7), Inches(11.2), Inches(3.8))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "PROJECT PRESENTATION  •  EXPLAINABLE AI & MULTIMODAL VERIFICATION"
    p0.font.size = Pt(11.5)
    p0.font.bold = True
    p0.font.color.rgb = CYAN
    p0.space_after = Pt(8)

    p1 = tf1.add_paragraph()
    p1.text = "VERITAS: Explainable Fake & AI News Detection"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(10)

    p2 = tf1.add_paragraph()
    p2.text = "From Static NLP Claim Benchmarks to Real-Time Cross-Source Corroboration"
    p2.font.size = Pt(16)
    p2.font.bold = True
    p2.font.color.rgb = EMERALD
    p2.space_after = Pt(10)

    p3 = tf1.add_paragraph()
    p3.text = "A multi-modal defense architecture that unifies real-time Google News RSS cross-checking, formal evidence mathematics, C2PA & SynthID synthetic media forensics, and zero-friction client intervention."
    p3.font.size = Pt(12)
    p3.font.color.rgb = TEXT_MUTED
    p3.space_after = Pt(14)

    p4 = tf1.add_paragraph()
    p4.text = "Key Inventions: Real-Time Corroboration  |  Evidence Score Math (E)  |  Media Provenance  |  Chrome MV3 & WhatsApp Bot"
    p4.font.size = Pt(11)
    p4.font.color.rgb = AMBER

    add_card(
        s1, left=0.8, top=5.6, width=11.733, height=1.35,
        title="",
        items=[
            ("The Fundamental Problem:", "Black-box AI models output opaque percentage scores that hallucinate and perpetuate publisher bias. Veritas replaces opacity with verifiable primary sources and mathematically bounded evidence."),
            ("System Breakthrough:", "Deterministic Evidence Score (E = 0.30M + 0.30C + 0.20F + 0.20S) delivered via in-place Shadow DOM browser banners and bilingual WhatsApp chat.")
        ],
        accent=EMERALD
    )

    # =========================================================================
    # SLIDE 2: BASE PAPER ANALYSIS & CRITICAL RESEARCH GAPS
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Base Paper: Methodology, Benchmark & Critical Gaps", "Literature Baseline", "Base Paper Analysis", ROSE, 2)

    add_card(
        s2, left=0.8, top=1.75, width=3.65, height=5.1,
        title="1. The Base Paper Benchmark",
        items=[
            ("Reference Paper:", "William Yang Wang (ACL 2017): '\"Liar, Liar Pants on Fire\": A New Benchmark Dataset for Fake News Detection'."),
            ("Dataset (LIAR):", "12,836 labeled short statements from PolitiFact rated across 6 fine-grained truth classes."),
            ("Model Methodology:", "Trained text classifiers (Logistic Regression, SVM, CNN, Bi-LSTM) on claim strings + metadata."),
            ("Baseline Finding:", "Adding metadata (speaker, party, credit history) boosted accuracy, but exposed severe real-world vulnerabilities.")
        ],
        accent=PURPLE
    )

    add_card(
        s2, left=4.85, top=1.75, width=3.65, height=5.1,
        title="2. Fatal Limitations of Base Paper",
        items=[
            ("[LIMITATION] Static Vacuum:", "Evaluates isolated statements offline. Never checks live web or credible news wires to verify ongoing events."),
            ("[LIMITATION] Metadata Bias:", "Memorizes speaker identity and party affiliation rather than evaluating factual claims. Fails completely on new publishers."),
            ("[LIMITATION] Black-Box Opacity:", "Outputs an arbitrary softmax probability. Cites zero primary source links or verifiable facts to readers."),
            ("[LIMITATION] Blind to Generative AI:", "Authored before modern generative LLMs, Midjourney, and synthetic deepfakes; lacks provenance detection.")
        ],
        accent=ROSE
    )

    add_card(
        s2, left=8.9, top=1.75, width=3.65, height=5.1,
        title="3. The Research Gaps Veritas Solves",
        items=[
            ("[GAP] Static to Real-Time:", "Misinformation spreads in minutes. Fact-checking demands active web retrieval rather than frozen model weights."),
            ("[GAP] Truth to Evidence:", "Algorithms cannot act as arbiters of absolute truth; they must surface consensus across trusted, independent newsrooms."),
            ("[GAP] Offline to In-Place:", "Academic benchmark scripts cannot protect real users. Production demands direct browser banners and chat alerts.")
        ],
        accent=AMBER
    )

    # =========================================================================
    # SLIDE 3: PROPOSED SYSTEM & WHAT CHANGES WE MADE
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Proposed System: How Veritas Changes & Advances the Baseline", "Proposed System", "Core Engineering Shifts", CYAN, 3)

    add_card(
        s3, left=0.8, top=1.75, width=3.65, height=5.1,
        title="1. Live Cross-Source Retrieval",
        items=[
            ("Base Paper Baseline:", "Classified isolated past sentences with zero external news lookup."),
            ("Veritas Engineering Shift:", "Asynchronous parallel Google News RSS retrieval engine across 5–10 reference outlets."),
            ("Tiered Source Registry:", "Weighted Tier 1 (Reuters, BBC, PTI: 1.00), Tier 2 (National Dailies: 0.85), Tier 3 (Regional: 0.65)."),
            ("Freshness Guardrails:", "Breaking news filter with median reference age check (<=2.0h) prevents falsely penalizing developing news.")
        ],
        accent=CYAN
    )

    add_card(
        s3, left=4.85, top=1.75, width=3.65, height=5.1,
        title="2. Formal Evidence Math (E)",
        items=[
            ("Base Paper Baseline:", "Opaque softmax probability from neural network black-boxes."),
            ("Veritas Engineering Shift:", "Deterministic Evidence Score bounded 0–100 with verifiable math:"),
            ("Mathematical Formula:", "E = 0.30·M + 0.30·C + 0.20·F + 0.20·S"),
            ("M (Semantic Match):", "Recall-focused content word overlap."),
            ("C (Consensus):", "Supporting vs. refuting outlet ratio."),
            ("F (Factual Consistency):", "Entities, dates & Google Fact Check API."),
            ("S (Source Quality):", "Journalistic credibility weighting."),
            ("Contradiction Invariant:", "-15 penalty if Tier-1 outlet refutes.")
        ],
        accent=EMERALD
    )

    add_card(
        s3, left=8.9, top=1.75, width=3.65, height=5.1,
        title="3. Multimodal Forensics & UI",
        items=[
            ("Base Paper Baseline:", "Text-only ASCII claim strings; zero user interface."),
            ("Veritas Engineering Shift:", "Multi-modal provenance + dual real-world delivery platforms:"),
            ("C2PA & SynthID Forensics:", "Decodes cryptographic JUMBF manifests & imperceptible DeepMind watermarks."),
            ("Stylometric AI Cues:", "Flags flat sentence burstiness, lexical TTR, and stock LLM clichés ('delve', 'tapestry')."),
            ("In-Place Chrome MV3:", "Shadow DOM banner with 0% style collision."),
            ("Bilingual WhatsApp Bot:", "Twilio webhook (Port 8001) in EN + Hindi.")
        ],
        accent=PURPLE
    )

    # =========================================================================
    # SLIDE 4: SYSTEM PIPELINE (COLOR-CODED: DONE VS NOT DONE)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "End-to-End Pipeline Architecture: Completed vs. Planned Work", "Pipeline Architecture", "Implementation Status", EMERALD, 4)

    # Top Banner Legend
    add_card(
        s4, left=0.8, top=1.65, width=11.733, height=0.72,
        title="",
        items=[
            ("PIPELINE STATUS COLOR CODING:  ", "[DONE] = Fully implemented, tested & operational (Emerald Green)    |    [NOT DONE] = In-progress or future architecture (Amber)")
        ],
        accent=None,
        bg=CARD_BG,
        border_color=CARD_BORDER
    )

    # STAGE 1: INGESTION & CLIENT UX
    add_card(
        s4, left=0.8, top=2.52, width=2.75, height=4.45,
        title="Stage 1: Client & Ingestion",
        items=[
            ("[DONE] Shadow DOM Banner:", "MV3 Chrome extension injects isolated Shadow DOM; 0% CSS clash."),
            ("[DONE] Service Worker Relay:", "MV3 background worker proxies requests, avoiding mixed-content."),
            ("[DONE] WhatsApp Webhook:", "Twilio endpoint (Port 8001) parses forwarded URLs & text."),
            ("[DONE] Anti-SSRF Shield:", "Blocks private IP ranges & validates HTTP/HTTPS schemes."),
            ("---", "---"),
            ("[NOT DONE] Safari & Firefox:", "Cross-browser port to Firefox MV2/MV3 and Safari WebExtension."),
            ("[NOT DONE] Mobile Share-Sheet:", "iOS/Android native share intent receiver for instant fact-checking.")
        ],
        accent=EMERALD,
        border_color=CARD_BORDER
    )

    # STAGE 2: ORCHESTRATION & RETRIEVAL
    add_card(
        s4, left=3.8, top=2.52, width=2.75, height=4.45,
        title="Stage 2: Cross-Checking",
        items=[
            ("[DONE] FastAPI Pipeline:", "Parallel sub-second orchestration across all 4 detection engines."),
            ("[DONE] Claim Distillation:", "Regex rule engine with Gemini structured claim fallback."),
            ("[DONE] Parallel RSS Feeds:", "Dispatches 2–3 targeted queries to Google News across Tier 1/2/3."),
            ("[DONE] Consensus Engine:", "Weighted consensus matrix computing supporting vs refuting pools."),
            ("---", "---"),
            ("[NOT DONE] Redis Cache:", "Distributed multi-region cache (currently thread-safe in-memory TTL)."),
            ("[NOT DONE] Async Queues:", "Celery/RabbitMQ worker queue for 1,000+ simultaneous requests.")
        ],
        accent=EMERALD,
        border_color=CARD_BORDER
    )

    # STAGE 3: AI FORENSICS & GATED ML
    add_card(
        s4, left=6.8, top=2.52, width=2.75, height=4.45,
        title="Stage 3: AI Forensics & ML",
        items=[
            ("[DONE] C2PA Provenance:", "Cryptographic JUMBF manifest & IPTC Digital Source validation."),
            ("[DONE] SynthID Markers:", "Detects DeepMind imperceptible watermarks & generator EXIF tags."),
            ("[DONE] Stylometrics:", "Sentence burstiness, lexical diversity (TTR) & stock LLM dictionary."),
            ("[DONE] 70% Precision Gate:", "Blocks unqualified ML models from altering verdicts (shadow mode)."),
            ("---", "---"),
            ("[NOT DONE] IndicBERT SLM:", "Fine-tuned multilingual transformer for regional Hindi & Marathi news."),
            ("[NOT DONE] Video Deepfakes:", "Frame-by-frame facial landmark jitter and audio spectral cloning.")
        ],
        accent=AMBER,
        border_color=CARD_BORDER
    )

    # STAGE 4: VERDICT & DELIVERY
    add_card(
        s4, left=9.8, top=2.52, width=2.733, height=4.45,
        title="Stage 4: Verdict & Output",
        items=[
            ("[DONE] Evidence Math E:", "Deterministic bounded score E = 0.30M + 0.30C + 0.20F + 0.20S."),
            ("[DONE] Contradiction Guard:", "Flat -15 penalty applied when trusted sources explicitly debunk."),
            ("[DONE] Verdict Hierarchy:", "Maps to TRUSTED (E>65), UNCERTAIN, SUSPICIOUS, or FAKE NEWS."),
            ("[DONE] Bilingual Verdicts:", "Formatted English & Hindi verdict cards returned to WhatsApp."),
            ("---", "---"),
            ("[NOT DONE] Crowd Feedback:", "User appeal portal & community journalist dispute mechanism.")
        ],
        accent=EMERALD,
        border_color=CARD_BORDER
    )

    # =========================================================================
    # SLIDE 5: PROJECT PROGRESS, EMPIRICAL METRICS & NEXT STEPS
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Project Progress, Empirical Verification & Phase Roadmap", "Project Status", "Milestones & Delivery", CYAN, 5)

    add_stat_box(s5, left=0.8, top=1.7, width=2.7, height=1.35, number="80%", label="Total Pipeline Built", sublabel="Core Engine, MV3 & Bot Live", accent=EMERALD)
    add_stat_box(s5, left=3.8, top=1.7, width=2.7, height=1.35, number="262", label="Automated Unit Tests", sublabel="100% Pass Rate in Pytest", accent=CYAN)
    add_stat_box(s5, left=6.8, top=1.7, width=2.7, height=1.35, number="< 1.8s", label="Analysis Latency", sublabel="End-to-End Corroboration", accent=PURPLE)
    add_stat_box(s5, left=9.8, top=1.7, width=2.733, height=1.35, number="2", label="Live Client Interfaces", sublabel="Chrome Extension & WhatsApp", accent=AMBER)

    add_card(
        s5, left=0.8, top=3.25, width=5.65, height=3.65,
        title="Completed Work (Phase 1 & Phase 2 Delivered)",
        items=[
            ("Core Verification Engine:", "FastAPI backend running locally with zero heavy dependencies, sub-second claim extraction, and anti-SSRF firewall."),
            ("Parallel Cross-Checking:", "Real-time Google News RSS query generator matching claims against Tier 1/2/3 wire services."),
            ("Multi-Modal Forensics:", "Image C2PA manifest decoding, IPTC synthetic tagging, DeepMind SynthID identification, and stylometric LLM burstiness."),
            ("Dual User Interfaces:", "Fully functional Manifest V3 Chrome Extension with Shadow DOM banner + Twilio WhatsApp Bot supporting bilingual English/Hindi output."),
            ("Production Robustness:", "262 offline deterministic unit tests + 1-hour chaos soak test script guaranteeing 9 system invariants.")
        ],
        accent=EMERALD
    )

    add_card(
        s5, left=6.85, top=3.25, width=5.68, height=3.65,
        title="Pending Work & Phase 3 Roadmap (Next Steps)",
        items=[
            ("Distributed Infrastructure:", "Transitioning in-memory RSS caches to a multi-node Redis cluster for production scaling."),
            ("Fine-Tuned Multilingual SLM:", "Training an IndicBERT model for deep contextual nuances in regional Hindi, Marathi, and Tamil news."),
            ("Real-Time Deepfake Video Analysis:", "Integrating video frame facial landmark jitter and audio clone frequency spectral analysis."),
            ("Multi-Browser Ecosystem:", "Porting the MV3 extension to Mozilla Firefox and Apple Safari desktop/iOS browsers."),
            ("Community Feedback Portal:", "Building a web-based public appeal dashboard where journalists can contest algorithmic flags.")
        ],
        accent=AMBER
    )

    prs.save(output_path)
    print(f"5-Slide Masterpiece Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_veritas_5slide_presentation()
